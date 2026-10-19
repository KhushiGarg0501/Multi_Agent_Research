import streamlit as st
from graph import graph

from rag import (
    index_pdf,
    answer_question,
    list_indexed_documents,
)
import io
from docx import Document
from fpdf import FPDF
from io import BytesIO
import re
from docx import Document
from docx.shared import Inches, Pt
from io import BytesIO
from xml.sax.saxutils import escape
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.units import inch
from reportlab.platypus import (
    SimpleDocTemplate,
    Paragraph,
    Spacer,
    ListFlowable,
    ListItem,
)


def create_report_pdf(report: str) -> bytes:
    """Create a formatted PDF from the research report."""
    buffer = BytesIO()

    doc = SimpleDocTemplate(
        buffer,
        pagesize=A4,
        rightMargin=0.7 * inch,
        leftMargin=0.7 * inch,
        topMargin=0.7 * inch,
        bottomMargin=0.7 * inch,
        title="Research Report",
        author="AI Research System",
    )

    base_styles = getSampleStyleSheet()

    title_style = ParagraphStyle(
        "ReportTitle",
        parent=base_styles["Title"],
        fontSize=20,
        leading=26,
        alignment=TA_CENTER,
        spaceAfter=20,
        textColor=colors.HexColor("#243B53"),
    )

    heading_style = ParagraphStyle(
        "ReportHeading",
        parent=base_styles["Heading1"],
        fontSize=15,
        leading=19,
        spaceBefore=14,
        spaceAfter=8,
        keepWithNext=True,
        textColor=colors.HexColor("#243B53"),
    )

    subheading_style = ParagraphStyle(
        "ReportSubheading",
        parent=base_styles["Heading2"],
        fontSize=12,
        leading=16,
        spaceBefore=10,
        spaceAfter=5,
        keepWithNext=True,
    )

    body_style = ParagraphStyle(
        "ReportBody",
        parent=base_styles["BodyText"],
        fontSize=9.5,
        leading=14,
        spaceAfter=7,
        splitLongWords=True,
    )

    story = []
    first_heading = True

    for raw_line in report.splitlines():
        line = raw_line.strip()

        if not line:
            story.append(Spacer(1, 4))
            continue

        # Markdown headings
        if line.startswith("# "):
            heading = line[2:].strip()

            if first_heading:
                story.append(
                    Paragraph(escape(heading), title_style)
                )
                first_heading = False
            else:
                story.append(
                    Paragraph(escape(heading), heading_style)
                )
            continue

        if line.startswith("## "):
            story.append(
                Paragraph(escape(line[3:].strip()), heading_style)
            )
            continue

        if line.startswith("### "):
            story.append(
                Paragraph(escape(line[4:].strip()), subheading_style)
            )
            continue

        # Basic Markdown bullet lists
        if line.startswith(("- ", "* ")):
            text = escape(line[2:].strip())
            story.append(
                ListFlowable(
                    [ListItem(Paragraph(text, body_style))],
                    bulletType="bullet",
                    leftIndent=18,
                )
            )
            story.append(Spacer(1, 4))
            continue

        # Basic numbered lists
        import re

        match = re.match(r"^\d+[.)]\s+(.+)$", line)
        if match:
            story.append(
                Paragraph(
                    escape(match.group(1)),
                    body_style,
                    bulletText="•",
                )
            )
            continue

        # Regular paragraphs
        text = escape(line)
        text = re.sub(r"\*\*(.+?)\*\*", r"<b>\1</b>", text)
        text = re.sub(r"(?<!\*)\*(.+?)\*(?!\*)", r"<i>\1</i>", text)

        story.append(Paragraph(text, body_style))

    if not story:
        story.append(Paragraph("No report content available.", body_style))

    def add_page_number(canvas, document):
        canvas.saveState()
        canvas.setFont("Helvetica", 8)
        canvas.setFillColor(colors.grey)
        canvas.drawCentredString(
            A4[0] / 2,
            0.35 * inch,
            f"Page {document.page}",
        )
        canvas.restoreState()

    doc.build(
        story,
        onFirstPage=add_page_number,
        onLaterPages=add_page_number,
    )

    return buffer.getvalue()
 
def create_report_docx(report: str) -> bytes:
    """Convert a Markdown-style research report into a Word document."""
    doc = Document()

    section = doc.sections[0]
    section.top_margin = Inches(0.75)
    section.bottom_margin = Inches(0.75)
    section.left_margin = Inches(0.85)
    section.right_margin = Inches(0.85)

    doc.styles["Normal"].font.name = "Arial"
    doc.styles["Normal"].font.size = Pt(10)

    for line in report.splitlines():
        line = line.strip()

        if not line:
            continue

        if line.startswith("# "):
            doc.add_heading(line[2:].strip(), level=0)

        elif line.startswith("## "):
            doc.add_heading(line[3:].strip(), level=1)

        elif line.startswith("### "):
            doc.add_heading(line[4:].strip(), level=2)

        elif line.startswith(("- ", "* ")):
            doc.add_paragraph(line[2:].strip(), style="List Bullet")

        elif re.match(r"^\d+\.\s+", line):
            text = re.sub(r"^\d+\.\s+", "", line)
            doc.add_paragraph(text, style="List Number")

        else:
            # Remove basic Markdown bold/italic markers.
            text = re.sub(r"\*\*(.*?)\*\*", r"\1", line)
            text = re.sub(r"(?<!\*)\*(.*?)\*(?!\*)", r"\1", text)
            doc.add_paragraph(text)

    buffer = BytesIO()
    doc.save(buffer)
    return buffer.getvalue()

# PAGE CONFIG


st.set_page_config(
    page_title="AI Research System",
    page_icon="🧠",
    layout="wide",
    initial_sidebar_state="expanded"
)


# CUSTOM CSS


st.markdown("""
<style>

    /* Main page */
    .stApp {
        background-color: #f8fafc;
    }

    /* Hide Streamlit default menu */
    #MainMenu {
        visibility: hidden;
    }

    footer {
        visibility: hidden;
    }

    /* Header */
    .hero {
        padding: 10px 0 25px 0;
    }

    .hero-title {
        font-size: 42px;
        font-weight: 800;
        color: #172033;
        margin-bottom: 4px;
    }

    .hero-subtitle {
        font-size: 18px;
        color: #64748b;
    }

    /* Status badge */
    .status-badge {
        display: inline-block;
        padding: 7px 14px;
        border-radius: 20px;
        background-color: #dcfce7;
        color: #166534;
        font-size: 14px;
        font-weight: 600;
    }

    /* Section headings */
    .section-title {
        font-size: 25px;
        font-weight: 700;
        color: #172033;
        margin-top: 10px;
        margin-bottom: 15px;
    }

    /* Agent cards */
    .agent-card {
        background: white;
        border: 1px solid #e2e8f0;
        border-radius: 14px;
        padding: 18px;
        text-align: center;
        min-height: 125px;
        box-shadow: 0 2px 8px rgba(15, 23, 42, 0.04);
    }

    .agent-icon {
        font-size: 30px;
        margin-bottom: 8px;
    }

    .agent-name {
        font-size: 16px;
        font-weight: 700;
        color: #1e293b;
    }

    .agent-status {
        font-size: 13px;
        color: #16a34a;
        margin-top: 5px;
    }

    /* Research box */
    .research-box {
        background: white;
        border: 1px solid #e2e8f0;
        border-radius: 16px;
        padding: 25px;
        box-shadow: 0 3px 12px rgba(15, 23, 42, 0.05);
    }

    /* Report box */
    .report-box {
        background: white;
        border: 1px solid #e2e8f0;
        border-radius: 16px;
        padding: 28px;
        margin-top: 15px;
        min-height: 200px;
    }

    /* Sidebar */
    section[data-testid="stSidebar"] {
        background-color: #eef2f7;
    }

    /* Buttons */
    .stButton > button {
        border-radius: 10px;
        font-weight: 600;
        min-height: 45px;
    }

    /* Divider */
    hr {
        border-color: #e2e8f0;
    }

</style>
""", unsafe_allow_html=True)


# ============================================================
# SIDEBAR
# ============================================================

with st.sidebar:

    st.markdown("## ⚙️ Research Settings")

    st.markdown(
        """
        <div style="
            background-color:white;
            padding:18px;
            border-radius:14px;
            border:1px solid #e2e8f0;
        ">
            <b>Multi-Agent Research</b><br><br>
            The system plans, searches, reads, writes,
            and critiques research automatically.
        </div>
        """,
        unsafe_allow_html=True
    )

    st.markdown("---")

    st.markdown("### 🤖 Research Agents")

    agents = [
        ("🧭", "Planner", "Breaks the topic into research questions"),
        ("🌐", "Search", "Finds relevant web sources"),
        ("📖", "Reader", "Extracts useful information"),
        ("✍️", "Writer", "Creates the research report"),
        ("🧐", "Critic", "Reviews report quality"),
    ]

    for icon, name, description in agents:
        st.markdown(
            f"""
            <div style="
                padding:10px 5px;
                margin-bottom:8px;
            ">
                <b>{icon} {name}</b><br>
                <span style="
                    font-size:12px;
                    color:#64748b;
                ">
                    {description}
                </span>
            </div>
            """,
            unsafe_allow_html=True
        )

    st.markdown("---")

    st.markdown("### 🛠️ Technology")

    st.caption("LangGraph")
    st.caption("Google Gemini")
    st.caption("Tavily Search")
    st.caption("BeautifulSoup")
    st.caption("Streamlit")


# ============================================================
# HEADER
# ============================================================

col1, col2 = st.columns([5, 1])

with col1:
    st.markdown(
        """
        <div class="hero">
            <div class="hero-title">🧠 AI Research System</div>
            <div class="hero-subtitle">
                Multi-Agent Research Assistant
            </div>
        </div>
        """,
        unsafe_allow_html=True
    )

with col2:
    st.markdown(
        """
        <div style="text-align:right; margin-top:15px;">
            <span class="status-badge">● System Ready</span>
        </div>
        """,
        unsafe_allow_html=True
    )

st.markdown(
    """
    <p style="color:#64748b; font-size:16px;">
        Generate structured research reports using AI agents,
        web research, source analysis, and critical review.
    </p>
    """,
    unsafe_allow_html=True
)

st.divider()


# ============================================================
# RESEARCH INPUT
# ============================================================

st.markdown(
    '<div class="section-title">🔍 Start Your Research</div>',
    unsafe_allow_html=True
)

topic = st.text_area(
    "Research Topic",
    placeholder="Example: Impact of Generative AI on Software Engineering",
    height=120,
    label_visibility="collapsed"
)

use_uploaded_docs = st.checkbox(
    "📚 Include evidence from my indexed PDFs",
    value=False,
    help=(
        "When enabled, the report writer retrieves relevant "
        "passages from your locally indexed PDF library."
    ),
)

col1, col2, col3 = st.columns([1.3, 1.3, 3])

with col1:
    research_button = st.button(
        "🚀 Start Research",
        use_container_width=True
    )

with col2:
    clear_button = st.button(
        "🗑️ Clear",
        use_container_width=True
    )

st.markdown('</div>', unsafe_allow_html=True)


# ============================================================
# AGENT DASHBOARD
# ============================================================

st.markdown(
    '<div class="section-title">🤖 Research Pipeline</div>',
    unsafe_allow_html=True
)

agent_columns = st.columns(5)

agent_data = [
    ("🧭", "Planner"),
    ("🌐", "Search"),
    ("📖", "Reader"),
    ("✍️", "Writer"),
    ("🧐", "Critic")
]

for column, (icon, name) in zip(agent_columns, agent_data):

    with column:

        st.markdown(
            f"""
            <div class="agent-card">
                <div class="agent-icon">{icon}</div>
                <div class="agent-name">{name}</div>
                <div class="agent-status">● Ready</div>
            </div>
            """,
            unsafe_allow_html=True
        )



# ============================================================
# RESEARCH WORKFLOW + PERSISTENT RESULTS
# ============================================================

if "research_result" not in st.session_state:
    st.session_state.research_result = None

if "research_topic" not in st.session_state:
    st.session_state.research_topic = ""

if clear_button:
    st.session_state.research_result = None
    st.session_state.research_topic = ""
    st.rerun()

if research_button:
    if not topic.strip():
        st.warning("Please enter a research topic first.")
    else:
        progress = st.progress(0)
        status = st.empty()

        try:
            status.info("Running research workflow...")

            progress.progress(10)

            result = graph.invoke({
                "topic": topic.strip(),
                "revision_count": 0,
                "use_uploaded_docs": use_uploaded_docs,
            })

            progress.progress(100)

            report = result.get("report", "")

            if not isinstance(report, str):
                report = str(report) if report else ""

            if not report.strip():
                st.error("The workflow finished but returned an empty report.")
            else:
                st.session_state.research_result = result
                st.session_state.research_topic = topic.strip()
                status.success("Research completed successfully!")

        except Exception as e:
            status.error("Research workflow failed.")
            st.exception(e)

# Render saved results OUTSIDE the button condition.
# This keeps the report visible after normal Streamlit reruns.

result = st.session_state.research_result

if result:
    report = result.get("report", "")
    report = report if isinstance(report, str) else str(report)

    st.divider()
    st.markdown("## 📄 Final Research Report")
    st.caption(f"Topic: {st.session_state.research_topic}")

    # No separate opening/closing HTML divs.
    # This removes the empty report container.
    st.markdown(report)

    st.divider()
    

# DOWNLOAD REPORT + RESEARCH DIAGNOSTICS

result = st.session_state.research_result

if result:
    report = result.get("report", "")
    report = report if isinstance(report, str) else str(report)

    if report.strip():
        st.subheader("📥 Download your report")

        col_pdf, col_docx = st.columns(2)

        with col_pdf:
            pdf_bytes = create_report_pdf(report)
            st.download_button(
                label="Download PDF",
                data=pdf_bytes,
                file_name="research_report.pdf",
                mime="application/pdf",
                use_container_width=True,
                key="download_report_pdf",
            )

        with col_docx:
            docx_bytes = create_report_docx(report)
            st.download_button(
                label="Download Word document (.docx)",
                data=docx_bytes,
                file_name="research_report.docx",
                mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                use_container_width=True,
                key="download_report_docx",
            )

    with st.expander(
        "🔍 Research Questions and Quality Review",
        expanded=False,
    ):
        st.markdown("### Research Questions")

        for i, question in enumerate(
            result.get("research_questions", []), 1
        ):
            st.markdown(f"{i}. {question}")

        st.markdown("### Critic Feedback")
        st.write(result.get("critique", "No critique returned."))

        if result.get("research_gaps"):
            st.markdown("### Research Gaps")
            st.write(result["research_gaps"])

        if result.get("fact_check"):
            st.markdown("### Fact-Check Findings")
            st.write(result["fact_check"])


# ============================================================
# PDF QUESTION-ANSWERING (RAG)
# ============================================================

st.divider()

st.markdown(
    '<div class="section-title">📚 Chat with PDFs</div>',
    unsafe_allow_html=True,
)

st.write(
    "Upload research papers or documents, index them, "
    "and ask questions using their contents."
)

uploaded_pdfs = st.file_uploader(
    "Upload PDF documents",
    type=["pdf"],
    accept_multiple_files=True,
    key="rag_pdf_uploader",
)

if st.button(
    "📥 Index uploaded PDFs",
    use_container_width=True,
    key="index_pdfs_button",
):
    if not uploaded_pdfs:
        st.warning("Please upload at least one PDF.")

    else:
        total_chunks = 0

        for pdf in uploaded_pdfs:
            try:
                with st.spinner(f"Indexing {pdf.name}..."):
                    result = index_pdf(
                        pdf.name,
                        pdf.getvalue(),
                    )

                total_chunks += result["chunks"]

                st.success(
                    f"Indexed {result['filename']}: "
                    f"{result['chunks']} text chunks."
                )

            except Exception as error:
                st.error(
                    f"Could not index {pdf.name}: {error}"
                )

        if total_chunks:
            st.success(
                f"Finished indexing {total_chunks} text chunks."
            )


# ------------------------------------------------------------
# SHOW DOCUMENTS ALREADY STORED LOCALLY
# ------------------------------------------------------------

try:
    indexed_documents = list_indexed_documents()

    with st.expander("📂 Indexed PDF library", expanded=False):
        if indexed_documents:
            for filename in indexed_documents:
                st.write(f"📄 {filename}")
        else:
            st.info("No PDFs have been indexed yet.")

except Exception as error:
    indexed_documents = []
    st.warning(f"Could not load the PDF library: {error}")


# ------------------------------------------------------------
# ASK QUESTIONS
# ------------------------------------------------------------

st.markdown("### 💬 Ask your documents")

pdf_question = st.text_input(
    "Your question",
    placeholder=(
        "Example: What are the main findings of these papers?"
    ),
    key="pdf_question_input",
)

if st.button(
    "🔎 Find answer",
    use_container_width=True,
    key="ask_pdf_question_button",
):
    if not pdf_question.strip():
        st.warning("Please enter a question.")

    elif not indexed_documents:
        st.warning(
            "Upload and index at least one PDF before asking questions."
        )

    else:
        try:
            with st.spinner(
                "Searching your PDFs and generating an answer..."
            ):
                answer, sources = answer_question(
                    pdf_question.strip(),
                    top_k=5,
                )

            st.session_state["rag_answer"] = answer
            st.session_state["rag_sources"] = sources
            st.session_state["rag_question"] = pdf_question.strip()

        except Exception as error:
            st.error("Could not answer the question.")
            st.code(str(error))


# Keep the last answer visible when Streamlit reruns.
if st.session_state.get("rag_answer"):
    st.markdown("### 📝 Answer")

    st.markdown(st.session_state["rag_answer"])

    with st.expander("📖 Retrieved source passages"):
        for source in st.session_state.get("rag_sources", []):
            st.write(
                f"**{source['filename']}** — "
                f"page {source['page']}"
            )


# ============================================================
# FOOTER
# ============================================================

st.divider()

st.markdown(
    """
    <div style="
        text-align:center;
        color:#94a3b8;
        padding:10px;
        font-size:13px;
    ">
        🧠 AI Research System &nbsp;•&nbsp;
        LangGraph &nbsp;•&nbsp;
        Gemini &nbsp;•&nbsp;
        Tavily &nbsp;•&nbsp;
        Streamlit
    </div>
    """,
    unsafe_allow_html=True
)